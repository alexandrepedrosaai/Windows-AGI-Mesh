#!/usr/bin/env python3
"""Gravador de AuditLog - M365 ChinCopilot.
Append-only: nunca sobrescreve, sempre adiciona nova linha na celula AuditLog do tenant.
Formato (por linha): YYYY-MM-DD HH:MM:SS (UTC) | agente | acao objeto | resultado | nota
Uso:
  python3 audit_log.py <intra-id-ou-query> --agent ChinCopilot --action "create user" \
      --object "john.doe@contoso.onmicrosoft.com" --result success --note "ticket #4521"
"""
import argparse, sys
from datetime import datetime, timezone
import openpyxl
from resolve_tenant import load_tenants, resolve

def append_audit(matrix, intra_id, agent, action, obj, result, note):
    wb = openpyxl.load_workbook(matrix)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c else "" for c in rows[0]]
    col = {n: i for i, n in enumerate(header) if n}
    if "AuditLog" not in col or "IntraID" not in col:
        print("ERRO: colunas IntraID/AuditLog faltando na matriz.", file=sys.stderr)
        sys.exit(1)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S (UTC)")
    entry = f"{ts} | {agent} | {action} {obj} | {result} | {note}"
    target_row = None
    for rnum, row in enumerate(rows[1:], start=2):
        val = row[col["IntraID"]]
        if val is not None and str(val).strip() == intra_id:
            target_row = rnum
            break
    if not target_row:
        print(f"ERRO: IntraID '{intra_id}' nao encontrado na matriz.", file=sys.stderr)
        sys.exit(1)
    cell = ws.cell(row=target_row, column=col["AuditLog"] + 1)
    existing = str(cell.value).strip() if cell.value else ""
    updated = (existing + "\n" + entry) if existing else entry
    if len(updated) > 32767:
        print("ERRO: AuditLog excede o limite de 32767 caracteres por celula; entrada nao gravada.", file=sys.stderr)
        sys.exit(1)
    cell.value = updated
    cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    wb.save(matrix)
    print(f"AUDIT OK (IntraID={intra_id}): {entry}")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("query")
    ap.add_argument("--matrix", default="_M365TenantintraID_MCPControl.xlsx")
    ap.add_argument("--agent", default="ChinCopilot")
    ap.add_argument("--action", required=True)
    ap.add_argument("--object", default="")
    ap.add_argument("--result", default="success", choices=["success","partial","failed"])
    ap.add_argument("--note", default="")
    ap.add_argument("--include-disabled", action="store_true")
    args = ap.parse_args()
    tenants = load_tenants(args.matrix, args.include_disabled)
    hits, tier = resolve(args.query, tenants)
    if not hits:
        print(f"ERRO: tenant '{args.query}' unmanaged - nao da pra auditar.", file=sys.stderr); sys.exit(1)
    if len(hits) > 1:
        print(f"ERRO: ambíguo ({len(hits)} candidatos). Especifique por IntraID.", file=sys.stderr); sys.exit(1)
    append_audit(args.matrix, hits[0]["IntraID"], args.agent, args.action, args.object, args.result, args.note)

if __name__ == "__main__":
    main()
