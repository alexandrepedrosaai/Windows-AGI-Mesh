#!/usr/bin/env python3
"""Constroi _M365TenantintraID_MCPControl.xlsx - M365 ChinCopilot.
Cabecalhos obrigatorios + uma linha de exemplo DESABILITADA (Enabled=no) + aba Instrucoes.
Nenhum dado real de tenant e inventado - a linha de exemplo e apenas formato.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HEADERS = ["IntraID","TenantId","PrimaryDomain","DisplayName","McpServerMode",
           "McpEndpoint","RuntimeVersion","Owner","Environment","Region",
           "Enabled","Scopes","Notes","AuditLog"]

OUT = "/home/user/.super_doubao/super-doubao-runtime/workspace/_M365TenantintraID_MCPControl.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Controle"

# Cabeçalho
hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF")
for c, h in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Linha de exemplo - DESABILITADA, nunca e alvo de operacoes
exemplo = [
    "EXEMPLO-0000",
    "00000000-0000-0000-0000-000000000000",
    "exemplo.onmicrosoft.com",
    "Tenant Exemplo (remover esta linha)",
    "stdio",
    "C:\\Mcp\\M365.Mcp.Server.dll",
    "10.0.13",
    "seu-time",
    "prod",
    "BR-South",
    "no",
    "user.read,group.readwrite",
    "Linha de formato - remover antes de usar",
    "",
]
for c, v in enumerate(exemplo, start=1):
    cell = ws.cell(row=2, column=c, value=v)
    cell.font = Font(italic=True, color="808080")

# Larguras
widths = [14,38,28,32,14,38,14,16,12,12,10,28,30,40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# Aba de instrucoes
ws2 = wb.create_sheet("Instrucoes")
instrucoes = [
    ["MATRIZ DE CONTROLE M365 / MCP - M365 ChinCopilot"],
    [""],
    ["1. Remova a linha EXEMPLO-0000 da aba Controle antes de usar."],
    ["2. Uma linha por tenant. Preencha TODAS as colunas obrigatorias:"],
    ["   IntraID (chave primaria, unica), TenantId (GUID), PrimaryDomain, DisplayName,"],
    ["   McpServerMode (stdio ou http), McpEndpoint (path .dll OU URL http/sse), RuntimeVersion (10.0.13)."],
    ["3. Enabled=no faz o tenant ser ignorado em operacoes (use para desativar sem apagar)."],
    ["4. Scopes restringe quais ferramentas MCP podem ser chamadas naquele tenant."],
    ["5. AuditLog e append-only: nunca sobrescreva, sempre adicione nova linha no campo."],
    ["   Formato: YYYY-MM-DD HH:MM:SS (UTC) | agente | acao objeto | resultado | nota"],
    [""],
    ["Validar: python3 scripts/validate_control_matrix.py _M365TenantintraID_MCPControl.xlsx --json"],
    ["Resolvedor de tenant: python3 resolve_tenant.py <intra-id-ou-dominio-ou-nome>"],
]
for r, row in enumerate(instrucoes, start=1):
    ws2.cell(row=r, column=1, value=row[0])
ws2.column_dimensions["A"].width = 100
ws2.cell(row=1, column=1).font = Font(bold=True, size=14)

wb.save(OUT)
print(f"Matriz criada: {OUT}")
